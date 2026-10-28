"""Tests for the Excel portfolio report."""

from pathlib import Path

from openpyxl import load_workbook

from portfolio_optimizer.reporting import create_excel_report


def test_excel_report_has_expected_sheets_and_rows(tmp_path):
    output = tmp_path / "portfolio_report.xlsx"

    result = create_excel_report(output)

    assert result.exists()
    assert result.stat().st_size > 0

    workbook = load_workbook(result, read_only=False, data_only=True)
    try:
        assert workbook.sheetnames == [
            "Portfolio Comparison", "Asset Returns", "Price History"
        ]

        comparison = workbook["Portfolio Comparison"]
        assert comparison.max_row == 4
        assert comparison.max_column == 4

        returns = workbook["Asset Returns"]
        prices = workbook["Price History"]
        assert returns.max_row == 60
        assert prices.max_row == 61

        assert comparison.freeze_panes == "A2"
        assert comparison.auto_filter.ref is not None
    finally:
        workbook.close()
