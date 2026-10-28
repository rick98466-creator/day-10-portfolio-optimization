"""Generate a formatted Excel portfolio analysis report."""

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from portfolio_optimizer.comparison import compare_allocations
from portfolio_optimizer.data_loader import load_prices
from portfolio_optimizer.returns import calculate_returns
from portfolio_optimizer.validation import validate_prices


def create_excel_report(
    output_path: str | Path = "reports/portfolio_analysis.xlsx",
) -> Path:
    """Generate the portfolio report and return its path."""
    prices = load_prices("data/raw/investment_prices.csv")
    validate_prices(prices)
    returns = calculate_returns(prices)
    comparison = compare_allocations(returns)

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        comparison.to_excel(
            writer, sheet_name="Portfolio Comparison", index=False
        )
        returns.to_excel(writer, sheet_name="Asset Returns", index=False)
        prices.to_excel(writer, sheet_name="Price History", index=False)

        workbook = writer.book

        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            worksheet.sheet_view.showGridLines = False

            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(
                    fill_type="solid", fgColor="1F4E78"
                )
                cell.alignment = Alignment(horizontal="center")

            for column_cells in worksheet.columns:
                max_length = max(
                    len(str(cell.value)) if cell.value is not None else 0
                    for cell in column_cells
                )
                column_letter = get_column_letter(column_cells[0].column)
                worksheet.column_dimensions[column_letter].width = min(
                    max(max_length + 2, 12), 30
                )

        comparison_sheet = workbook["Portfolio Comparison"]
        for row in range(2, comparison_sheet.max_row + 1):
            comparison_sheet.cell(row, 2).number_format = "0.00%"
            comparison_sheet.cell(row, 3).number_format = "0.00"
            comparison_sheet.cell(row, 4).number_format = "0.00"

        for sheet_name in ("Asset Returns", "Price History"):
            worksheet = workbook[sheet_name]
            worksheet["A1"].alignment = Alignment(horizontal="center")

            for row in range(2, worksheet.max_row + 1):
                worksheet.cell(row, 1).number_format = "yyyy-mm-dd"
                for col in range(2, worksheet.max_column + 1):
                    worksheet.cell(row, col).number_format = (
                        "0.00%" if sheet_name == "Asset Returns" else "0.000000"
                    )

    return output


if __name__ == "__main__":
    report_path = create_excel_report()
    print(f"Excel report created: {report_path.resolve()}")
    print(f"File size: {report_path.stat().st_size:,} bytes")
